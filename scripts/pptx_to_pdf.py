from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor

pptx_path = 'presentations/Sistema-TDR-Presentacion.pptx'
out_pdf = 'presentations/Sistema-TDR-Presentacion.pdf'

prs = Presentation(pptx_path)
# Define slide size in points (pptx uses inches as we set earlier)
PAGE_WIDTH = 13.33 * inch
PAGE_HEIGHT = 7.5 * inch

c = canvas.Canvas(out_pdf, pagesize=(PAGE_WIDTH, PAGE_HEIGHT))

for slide in prs.slides:
    # Background: try to use a light background
    c.setFillColor(HexColor('#e9ebef'))
    c.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, stroke=0, fill=1)

    # Title
    title = ''
    try:
        if slide.shapes.title:
            title = slide.shapes.title.text
    except Exception:
        title = ''

    c.setFont('Helvetica-Bold', 34)
    c.setFillColor(HexColor('#030213'))
    c.drawString(0.7*inch, PAGE_HEIGHT - 1*inch, title)

    # Body (placeholders)
    body_text = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape == slide.shapes.title:
            continue
        text = shape.text
        if text and text.strip():
            body_text.append(text)

    c.setFont('Helvetica', 18)
    text_y = PAGE_HEIGHT - 1.6*inch
    for block in body_text:
        lines = block.split('\n')
        for ln in lines:
            if text_y < 0.7*inch:
                c.showPage()
                c.setFillColor(HexColor('#e9ebef'))
                c.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, stroke=0, fill=1)
                c.setFont('Helvetica', 18)
                text_y = PAGE_HEIGHT - 1*inch
            c.setFillColor(HexColor('#141414'))
            c.drawString(0.8*inch, text_y, ln)
            text_y -= 0.4*inch
        text_y -= 0.1*inch

    # Notes (optional) - draw at bottom small
    notes = ''
    try:
        notes = slide.notes_slide.notes_text_frame.text
    except Exception:
        notes = ''
    if notes:
        c.setFont('Helvetica-Oblique', 10)
        c.setFillColor(HexColor('#333333'))
        c.drawString(0.7*inch, 0.5*inch, 'Notas: ' + notes)

    c.showPage()

c.save()
print('PDF generado en', out_pdf)
