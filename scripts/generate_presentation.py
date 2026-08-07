from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Paleta basada en src/styles/globals.css (adaptada)
PRIMARY = (3, 2, 19)       # #030213
ACCENT = (233, 235, 239)   # #e9ebef (acento claro)
HIGHLIGHT = (212, 24, 61)  # #d4183d (destructive/acento rojo)
MUTED = (236, 236, 240)    # #ececf0
TEXT = (20, 20, 24)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# helper
def set_background(slide, rgb):
    r, g, b = rgb
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(r, g, b)


def add_title_slide(title, subtitle, note):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_background(slide, ACCENT)
    title_tf = slide.shapes.title
    title_tf.text = title
    title_tf.text_frame.paragraphs[0].font.size = Pt(40)
    title_tf.text_frame.paragraphs[0].font.bold = True
    title_tf.text_frame.paragraphs[0].font.color.rgb = RGBColor(*PRIMARY)

    subtitle_shape = slide.placeholders[1]
    subtitle_shape.text = subtitle
    subtitle_shape.text_frame.paragraphs[0].font.size = Pt(18)
    subtitle_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(*TEXT)
    slide.notes_slide.notes_text_frame.text = note


def add_content_slide(title, bullets, note, bg='light'):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_background(slide, ACCENT if bg=='light' else PRIMARY)
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(28)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(*PRIMARY)

    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for b in bullets:
        p = body.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*TEXT)

    slide.notes_slide.notes_text_frame.text = note


# Slides content
add_title_slide('Sistema de Gestión de TDR', 'Simplifica creación, revisión y archivo de documentos', 'Objetivo: presentar beneficios y próximos pasos al cliente.')

add_content_slide('¿Qué problema resolvemos?', [
    'Duplicidad de documentos y versiones descontroladas.',
    'Falta de trazabilidad: quién y cuándo cambió.',
    'Demoras en aprobaciones y comunicaciones internas.'
], 'Ejemplo: retrasos por buscar archivos en distintas carpetas.')

add_content_slide('¿Qué es un TDR?', [
    'Documento con tareas, responsabilidades y plazos.',
    'Se usa en contrataciones y acuerdos con proveedores/locadores.',
    'Clave para cumplimiento y control administrativo.'
], 'Explicar en términos sencillos: es la "receta" del trabajo a contratar.')

add_content_slide('Objetivos para usted', [
    'Centralizar documentos y evitar duplicados.',
    'Reducir tiempo de aprobación y notificaciones automáticas.',
    'Mantener registro claro de cambios para auditoría.'
], 'Enfocado en ahorro de tiempo y reducción de riesgos.')

add_content_slide('Cómo funciona (flujo simple)', [
    'Crear TDR → asignar participantes → notificar → revisar → aprobar → archivar.',
    'Cada acción queda registrada y visible para los autorizados.'
], 'Mostrar flujo en la demo práctica.')

add_content_slide('Roles principales', [
    'Administrador: configura el sistema y gestiona usuarios.',
    'Contratante: crea TDR y solicita aprobaciones.',
    'Locador/Proveedor: revisa y responde.',
    'Auditor: accede al historial y reportes.'
], 'Dar ejemplos de personas reales del cliente.')

add_content_slide('Beneficios para su institución', [
    'Menos errores y pérdida de documentos.',
    'Claridad en responsabilidades y tiempos.',
    'Evidencia para auditorías y cumplimiento.'
], 'Enfatizar beneficios económicos y operativos.')

add_content_slide('Seguridad y confiabilidad', [
    'Accesos por rol: control y permisos.',
    'Registro de cambios: quién, cuándo y qué.',
    'Recomendación: backups periódicos y SSL.'
], 'Evitar tecnicismos; enfocar en protección de información.')

add_content_slide('Integración y despliegue', [
    'Se integra con email para notificaciones.',
    'Despliegue progresivo: piloto con 1 equipo.',
    'Opcional: firma electrónica y almacenamiento externo.'
], 'Sugerir plan de 2–4 semanas para PoC.')

add_content_slide('Próximos pasos', [
    'Validación con usuarios clave (prueba piloto).',
    'Ajuste de roles y permisos según su estructura.',
    'Plan de despliegue y capacitación breve.'
], 'Proponer fecha para demo y POC.')

# save
out_path = 'presentations/Sistema-TDR-Presentacion.pptx'
import os
os.makedirs(os.path.dirname(out_path), exist_ok=True)
prs.save(out_path)
print('Saved to', out_path)
